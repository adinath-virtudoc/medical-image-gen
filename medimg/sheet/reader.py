"""Read, import and locate worksheets. Never writes to the user's original file."""
from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from medimg.errors import MedimgError

from . import _compat
from .columns import REQUIRED, STATUSES, norm_header, system_name


class SheetError(MedimgError):
    """A sheet problem with a cause and a fix."""


@dataclass
class Row:
    idx: int
    subject: str
    unit: str
    lecture: str
    prompt: str
    image_name: str
    status: str
    size: tuple[int, int] | None
    seed: int | None
    error: str
    rel_path: str | None
    system: dict[str, object]
    parse_error: str = ""          # reader-detected problem (bad size, bad seed, bad status)


@dataclass
class Problem:
    row: int
    status: str
    reason: str
    fix: str


@dataclass
class SheetReport:
    sheet_id: str
    original_name: str
    working_path: Path
    worksheet: str
    rows: list[Row]
    counts: dict[str, int]
    problems: list[Problem]
    filters: dict[str, list[str]] = field(default_factory=dict)


_SIZE_RE = re.compile(r"^\s*(\d+)\s*[xX×]\s*(\d+)\s*$")


def import_sheet(src: Path, output_root: Path) -> Path:
    src = Path(src)
    if not src.is_file():
        raise SheetError(f"Sheet file not found: {src}.", "Choose an existing .xlsx file.")
    sheets = Path(output_root) / "_sheets"
    sheets.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = sheets / f"{stamp}_{src.name}"
    n = 1
    while dest.exists():
        dest = sheets / f"{stamp}-{n}_{src.name}"
        n += 1
    shutil.copyfile(src, dest)
    pointer = sheets / "current.json"
    tmp = pointer.with_suffix(".json.tmp")
    tmp.write_text(json.dumps({"path": dest.name, "original_name": src.name}), encoding="utf-8")
    import os
    os.replace(tmp, pointer)
    return dest


def current_sheet(output_root: Path) -> Path | None:
    pointer = Path(output_root) / "_sheets" / "current.json"
    try:
        data = json.loads(pointer.read_text(encoding="utf-8"))
        p = Path(output_root) / "_sheets" / data["path"]
    except (OSError, ValueError, KeyError, TypeError):
        return None
    return p if p.is_file() else None


def _header_map(ws) -> dict[str, int]:
    out: dict[str, int] = {}
    for cell in ws[1]:
        key = norm_header(cell.value)
        if key and key not in out:
            out[key] = cell.column
    return out


def find_worksheet(wb):
    for ws in wb.worksheets:
        if all(h in _header_map(ws) for h in REQUIRED):
            return ws
    raise SheetError(
        "No worksheet has the required columns in its first row.",
        "Use the template: Subject, Unit, Lecture, Image Prompt, Image Name, Status.")


def _text(v: object) -> str:
    return "" if v is None else str(v).strip()


def read_rows(path: Path) -> tuple[str, list[Row]]:
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:  # corrupt / not xlsx
        raise SheetError(f"Cannot open the sheet: {e}", "Save it as .xlsx and try again.") from e
    try:
        ws = find_worksheet(wb)
        hmap = _header_map(ws)
        sys_cols = {system_name(k): c for k, c in hmap.items() if system_name(k)}
        rows: list[Row] = []
        for r in range(2, ws.max_row + 1):
            def cell(h: str):
                c = hmap.get(h)
                return ws.cell(row=r, column=c).value if c else None
            subject, unit, lecture = (_text(cell(h)) for h in ("subject", "unit", "lecture"))
            prompt, name = _text(cell("image prompt")), _text(cell("image name"))
            raw_status = _text(cell("status")).casefold()
            size_raw, seed_raw = _text(cell("size")), cell("seed")
            if not (subject or unit or lecture or prompt or name or size_raw):
                continue                              # blank row
            perr = ""
            status = raw_status or "pending"
            if status not in STATUSES:
                perr = f"Status must be one of {', '.join(STATUSES)}"
                status = "pending"
            size = None
            if size_raw:
                m = _SIZE_RE.match(size_raw)
                if m:
                    size = (int(m.group(1)), int(m.group(2)))
                elif not perr:
                    perr = "Size must look like 768x1344"
            seed = None
            if _text(seed_raw):
                try:
                    f = float(_text(seed_raw))
                    if f != int(f) or f < 0:
                        raise ValueError
                    seed = int(f)
                except ValueError:
                    perr = perr or "Seed must be a whole number"
            system = {h: ws.cell(row=r, column=c).value for h, c in sys_cols.items()
                      if ws.cell(row=r, column=c).value not in (None, "")}
            rel = (_compat.rel_path(subject, unit, lecture, name)
                   if subject and unit and lecture and name else None)
            rows.append(Row(idx=r, subject=subject, unit=unit, lecture=lecture, prompt=prompt,
                            image_name=name, status=status, size=size, seed=seed,
                            error=_text(system.get("Error")), rel_path=rel, system=system,
                            parse_error=perr))
        return ws.title, rows
    finally:
        wb.close()
