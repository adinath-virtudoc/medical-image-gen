"""Atomic sheet writes: load, modify, save to a temp file in the same directory, os.replace."""
from __future__ import annotations

import os
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from .columns import DISPLAY, SYSTEM, norm_header
from .reader import _header_map, find_worksheet


def _ensure(ws) -> dict[str, int]:
    hmap = _header_map(ws)
    last = max(hmap.values())
    template = ws.cell(row=1, column=last)
    for name in SYSTEM:
        if norm_header(name) not in hmap:
            last += 1
            c = ws.cell(row=1, column=last, value=name)
            if template.has_style:
                c.font, c.fill, c.border = copy(template.font), copy(template.fill), copy(template.border)
                c.alignment = copy(template.alignment)
            ws.column_dimensions[c.column_letter].width = max(16, len(name) + 4)
            hmap[norm_header(name)] = last
    return hmap


def _display(hmap: dict[str, int], ws) -> dict[str, int]:
    return {(DISPLAY.get(k) or next((s for s in SYSTEM if norm_header(s) == k), k)): c
            for k, c in hmap.items()}


def _save_atomic(wb, path: Path) -> None:
    fd, tmp = tempfile.mkstemp(dir=str(path.parent), prefix=path.name + ".", suffix=".tmp")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def _clean(v: object) -> object:
    if isinstance(v, Path):
        return str(v)
    return v


def _load(path: Path):
    """Load for editing, freezing formula cells to their cached values.

    openpyxl drops cached formula results on save, and the reader uses data_only=True,
    so a surviving formula would read as empty after the first write.
    """
    wb = load_workbook(path)
    try:
        if any(c.data_type == "f" for ws in wb.worksheets for row in ws.iter_rows() for c in row):
            cached = load_workbook(path, data_only=True)
            try:
                for ws in wb.worksheets:
                    cws = cached[ws.title]
                    for row in ws.iter_rows():
                        for c in row:
                            if c.data_type == "f":
                                c.value = cws[c.coordinate].value
            finally:
                cached.close()
    except BaseException:
        wb.close()
        raise
    return wb


def ensure_system_columns(path: Path) -> dict[str, int]:
    path = Path(path)
    wb = _load(path)
    try:
        ws = find_worksheet(wb)
        before = len(_header_map(ws))
        hmap = _ensure(ws)
        if len(hmap) != before:
            _save_atomic(wb, path)
        return _display(hmap, ws)
    finally:
        wb.close()


def write_rows(path: Path, updates: dict[int, dict[str, object]]) -> None:
    path = Path(path)
    wb = _load(path)
    try:
        ws = find_worksheet(wb)
        hmap = _ensure(ws)
        for idx, upd in updates.items():
            for header, value in upd.items():
                key = norm_header(header)
                if key not in hmap:
                    col = max(hmap.values()) + 1
                    ws.cell(row=1, column=col, value=header)
                    hmap[key] = col
                ws.cell(row=idx, column=hmap[key]).value = _clean(value)
        _save_atomic(wb, path)
    finally:
        wb.close()


def write_row(path: Path, idx: int, updates: dict[str, object]) -> None:
    write_rows(path, {idx: updates})
