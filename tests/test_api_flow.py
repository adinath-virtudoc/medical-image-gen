"""API: full flow with FakeEngine (upload -> batch -> gallery -> sheet -> download). Also shared helpers."""
from __future__ import annotations

import io
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from medimg import config
from medimg.assess import fit
from medimg.engines.fake_engine import FakeEngine
from medimg.server import create_app

REPO = Path(__file__).resolve().parent.parent
SAMPLE = REPO / "templates" / "image_prompts_sample.xlsx"
HEAD = ["Subject", "Unit", "Lecture", "Image Prompt", "Image Name", "Status"]
OK_PROMPT = "a 45-year-old adult man, clinical photo"


def make_client(monkeypatch, tmp_path, *, root=None, engine=None, profile=None, steps=4):
    """Build (client, app, engine, root) with every path redirected into tmp_path."""
    monkeypatch.setattr(config, "PROFILE_PATH", tmp_path / "device_profile.json")
    monkeypatch.setattr(config, "SETTINGS_PATH", tmp_path / "cfg" / "config.json")
    root = Path(root) if root else tmp_path / "root"
    engine = engine or FakeEngine(step_s=0.002)
    settings = config.Settings(output_root=root, steps=steps, unload_after_job=False)
    app = create_app(engine, settings, profile)
    return TestClient(app, base_url=f"http://127.0.0.1:{settings.port}", headers={"X-Medimg": "1"}), app, engine, root


def build_sheet(path, rows, extra_cols=()):
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD + list(extra_cols))
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def simple_rows(n, status="pending"):
    return [["Anat", "Hand", "L1", OK_PROMPT, f"img{i}", status] for i in range(n)]


def upload(client, path, name=None, **form):
    with open(path, "rb") as f:
        return client.post("/api/sheet", files={"file": (name or Path(path).name, f.read(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                           data=form)


def wait_job(client, job_id=None, timeout=20):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        s = client.get("/api/status").json()
        if s["job"] is None and s["last_job"] and (job_id is None or s["last_job"]["id"] == job_id):
            return s
        time.sleep(0.01)
    raise AssertionError("job did not finish in time")


def wait_for(pred, timeout=20):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        v = pred()
        if v:
            return v
        time.sleep(0.005)
    raise AssertionError("condition not reached in time")


def sheet_rows_from_bytes(data: bytes):
    ws = load_workbook(io.BytesIO(data)).worksheets[0]
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


def profile_dict():
    pts = [{"width": 512, "height": 512, "steps": 4, "total_s": 30.0, "per_step_s": 5.5, "overhead_s": 8.0},
           {"width": 1024, "height": 1024, "steps": 2, "total_s": 62.0, "per_step_s": 23.0, "overhead_s": 16.0}]
    return {"schema_version": 1, "created_at": "x", "app_version": "0", "backend": {"name": "mflux"},
            "machine": {"hostname": "testhost", "apple_chip": "Apple Test", "ram_total_gb": 16.0},
            "benchmark": {"load_s": 1.0, "points": pts, "fit": fit(pts)}, "actuals": [], "notes": []}


def test_full_flow(monkeypatch, tmp_path):
    client, app, engine, root = make_client(monkeypatch, tmp_path, profile=profile_dict())
    with client:
        r = upload(client, SAMPLE)
        assert r.status_code == 200, r.text
        rep = r.json()
        assert rep["counts"] == {"total": 3, "done": 0, "pending": 3, "blocked": 0, "failed": 0, "invalid": 0}
        assert rep["original_name"] == "image_prompts_sample.xlsx"
        assert rep["problems"] == [] and 1 <= rep["suggested_batch"] <= 3
        assert rep["estimate_per_image_s"] > 0 and "rows" not in rep
        assert rep["filters"]["subjects"]

        r = client.post("/api/jobs/batch", json={"count": 2, "steps": 4, "default_size": "512x512"})
        assert r.status_code == 202
        body = r.json()
        assert len(body["rows"]) == 2
        end = wait_job(client, body["job_id"])
        assert end["last_job"]["outcome"] == "done"

        imgs = client.get("/api/gallery/images").json()
        assert len(imgs) == 2
        for im in imgs:
            assert set(im) == {"id", "rel_path", "image_name", "url", "thumb_url", "record"}
            png = client.get(im["url"])
            assert png.status_code == 200 and png.content[:8] == b"\x89PNG\r\n\x1a\n"
            th = client.get(im["thumb_url"])
            assert th.status_code == 200 and th.headers["content-type"] == "image/jpeg"
            assert client.get(f"/api/images/{im['id']}").json()["id"] == im["id"]
            assert (root / im["rel_path"]).exists()
        tree = client.get("/api/gallery/tree").json()
        assert tree["subjects"] and tree["single_count"] == 0

        rep2 = client.get("/api/sheet").json()
        assert rep2["counts"]["done"] == 2 and rep2["counts"]["pending"] == 1

        dl = client.get("/api/sheet/download")
        assert dl.status_code == 200
        assert 'filename="image_prompts_sample.xlsx"' in dl.headers["content-disposition"]
        rows = sheet_rows_from_bytes(dl.content)
        done = [r for r in rows if r["Status"] == "done"]
        assert len(done) == 2
        for r in done:
            assert r["Seed Used"] is not None and r["Seed"] == r["Seed Used"]
            assert r["Generation Time (s)"] is not None and r["Machine"] == "testhost · Apple Test"
            assert r["Model"] == "fake/fake-model" and r["Output Path"] and r["Steps"] == 4
        assert [r["Status"] for r in rows].count("pending") == 1


def test_original_file_never_modified(monkeypatch, tmp_path):
    src = build_sheet(tmp_path / "mine.xlsx", simple_rows(2))
    before = src.read_bytes()
    client, *_ = make_client(monkeypatch, tmp_path)
    with client:
        assert upload(client, src).status_code == 200
        client.post("/api/jobs/batch", json={"count": 2, "steps": 4})
        wait_job(client)
    assert src.read_bytes() == before
