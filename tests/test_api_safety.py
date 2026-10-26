"""Acceptance 3 and 6: the adults-only gate on every path, path-traversal 404s, loopback-only bind."""
from __future__ import annotations

import dataclasses
import json
import re
from pathlib import Path

from openpyxl import Workbook

from medimg import rules, sheet
from tests.test_api_flow import (HEAD, OK_PROMPT, build_sheet, make_client, sheet_rows_from_bytes,
                                 simple_rows, upload, wait_job)

def raw_status(app, raw_path: bytes) -> int:
    """GET with an exact raw path through the ASGI app (TestClient would normalise '..%2F' away)."""
    import asyncio
    from urllib.parse import unquote

    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": "GET",
             "scheme": "http", "path": unquote(raw_path.decode()), "raw_path": raw_path, "query_string": b"",
             "headers": [(b"host", b"127.0.0.1:8750")], "client": ("127.0.0.1", 1), "server": ("testserver", 80)}
    out = {}

    async def receive():
        return {"type": "http.request", "body": b"", "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            out["status"] = msg["status"]

    asyncio.run(app(scope, receive, send))
    return out["status"]


SERVER_SRC = Path(__file__).resolve().parent.parent / "medimg" / "server.py"


def mixed_sheet(path):
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD + ["Size"])
    ws.append(["Anat", "Hand", "L1", OK_PROMPT, "done_one", "done", ""])                      # 2 done
    ws.append(["Anat", "Hand", "L1", OK_PROMPT, "good", "", ""])                              # 3 pending
    ws.append(["Anat", "Hand", "L1", "a 10-year-old child with a cast", "minor", "", ""])     # 4 blocked
    ws.append(["Anat", "Hand", "L1", "", "noprompt", "", ""])                                 # 5 invalid
    ws.append(["Anat", "Hand", "L1", OK_PROMPT, "dup", "", ""])                               # 6 invalid
    ws.append(["Anat", "Hand", "L1", OK_PROMPT, "dup", "", ""])                               # 7 invalid
    ws.append(["Anat", "Hand", "L1", OK_PROMPT, "badsize", "", "huge"])                       # 8 invalid
    wb.save(path)
    return path


def test_acceptance_3_mixed_sheet(monkeypatch, tmp_path):
    client, _, engine, root = make_client(monkeypatch, tmp_path)
    with client:
        rep = upload(client, mixed_sheet(tmp_path / "mixed.xlsx")).json()
        assert rep["counts"] == {"total": 7, "done": 1, "pending": 1, "blocked": 1, "failed": 0, "invalid": 4}
        by_row = {p["row"]: p for p in rep["problems"]}
        assert set(by_row) == {4, 5, 6, 7, 8}
        assert by_row[4]["status"] == "blocked" and by_row[4]["reason"].startswith("Blocked:")
        assert "Missing Image Prompt" in by_row[5]["reason"]
        assert "Same output path as row 7" in by_row[6]["reason"] and "row 6" in by_row[7]["reason"]
        assert "Size must look like" in by_row[8]["reason"]
        assert all(p["fix"] for p in rep["problems"])

        rows = sheet_rows_from_bytes(client.get("/api/sheet/download").content)   # reasons persisted
        assert [r["Status"] or "pending" for r in rows] == ["done", "pending", "blocked", "invalid",
                                                            "invalid", "invalid", "invalid"]
        assert rows[2]["Error"].startswith("Blocked:")

        r = client.post("/api/jobs/batch", json={"count": 50, "steps": 4})
        assert r.status_code == 202 and r.json()["rows"] == [3]
        wait_job(client, r.json()["job_id"])
        assert engine.calls == 1
        assert [i["image_name"] for i in client.get("/api/gallery/images").json()] == ["good"]


def test_single_blocked_422(monkeypatch, tmp_path):
    client, _, engine, _ = make_client(monkeypatch, tmp_path)
    with client:
        for prompt in ("a teen boy with a sprained wrist", "a 16 yo patient", "a woman with a rash",
                       "a 17-year-old adult"):
            r = client.post("/api/jobs/single", json={"prompt": prompt, "steps": 4, "width": 512, "height": 512})
            assert r.status_code == 422, prompt
            err = r.json()["error"]
            assert err["cause"] == rules.check_prompt(prompt) and err["fix"] == rules.HINT
        assert engine.calls == 0 and client.get("/api/status").json()["job"] is None
        assert client.get("/api/status").json()["last_job"] is None


def test_regenerate_blocked_after_sidecar_tamper(monkeypatch, tmp_path):
    client, _, engine, root = make_client(monkeypatch, tmp_path)
    with client:
        r = client.post("/api/jobs/single", json={"prompt": OK_PROMPT, "steps": 4, "width": 512,
                                                   "height": 512, "image_name": "victim"})
        wait_job(client, r.json()["job_id"])
        img = client.get("/api/gallery/images").json()[0]
        side = root / "_single" / "victim.json"
        data = json.loads(side.read_text())
        data["prompt"] = "a young girl, clinical photo"
        side.write_text(json.dumps(data))
        png = root / "_single" / "victim.png"
        before = png.read_bytes()
        calls = engine.calls
        r = client.post("/api/jobs/regenerate", json={"image_id": img["id"]})
        assert r.status_code == 422
        assert r.json()["error"]["cause"].startswith("Blocked:") and r.json()["error"]["fix"] == rules.HINT
        assert engine.calls == calls and png.read_bytes() == before


def test_regenerate_ok_updates_sheet_row(monkeypatch, tmp_path):
    src = build_sheet(tmp_path / "work.xlsx", simple_rows(1))
    client, _, engine, root = make_client(monkeypatch, tmp_path)
    with client:
        upload(client, src)
        wait_job(client, client.post("/api/jobs/batch", json={"count": 1, "steps": 4}).json()["job_id"])
        img = client.get("/api/gallery/images").json()[0]
        old_seed = img["record"]["seed"]
        r = client.post("/api/jobs/regenerate", json={"image_id": img["id"]})
        assert r.status_code == 202 and r.json()["sheet_row"] == 2
        new_seed = r.json()["seed"]
        wait_job(client, r.json()["job_id"])
        assert client.get(f"/api/images/{img['id']}").json()["seed"] == new_seed
        rows = sheet_rows_from_bytes(client.get("/api/sheet/download").content)
        assert rows[0]["Seed Used"] == new_seed and rows[0]["Status"] == "done"
        assert len(client.get("/api/gallery/images").json()) == 1
        assert old_seed is not None


def test_batch_rows_blocked_by_server_gate(monkeypatch, tmp_path):
    """Even if validation drifted and left a minor-term row pending, the route blocks it."""
    rows = simple_rows(2)
    rows[0][3] = "a young boy with a fracture"
    src = build_sheet(tmp_path / "work.xlsx", rows)
    client, _, engine, root = make_client(monkeypatch, tmp_path)
    monkeypatch.setattr(sheet, "validate", lambda rs, **kw: [dataclasses.replace(r, status="pending", error="")
                                                             for r in rs])
    with client:
        upload(client, src)
        r = client.post("/api/jobs/batch", json={"count": 5, "steps": 4})
        assert r.status_code == 202 and r.json()["rows"] == [3]
        wait_job(client, r.json()["job_id"])
        assert engine.calls == 1
        out = sheet_rows_from_bytes(client.get("/api/sheet/download").content)
        assert out[0]["Status"] == "blocked" and out[0]["Error"].startswith("Blocked:")
        assert out[1]["Status"] == "done"
        assert not (root / "Anat/Hand/L1/img0.png").exists()
        # the stored status is now blocked; with the real validator it stays blocked and is not generated
        monkeypatch.undo()
        again = client.post("/api/jobs/batch", json={"count": 5, "steps": 4})
        assert again.status_code == 400


def test_edited_working_sheet_row_blocked_on_batch(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    src = build_sheet(tmp_path / "work.xlsx", simple_rows(1))
    client, app, engine, root = make_client(monkeypatch, tmp_path)
    with client:
        upload(client, src)
        path = app.state.sheet_ref.path
        wb = load_workbook(path)
        wb.worksheets[0]["D2"] = "a toddler"
        wb.save(path)
        r = client.post("/api/jobs/batch", json={"count": 1, "steps": 4})
        assert r.status_code == 400 and engine.calls == 0
        assert sheet_rows_from_bytes(client.get("/api/sheet/download").content)[0]["Status"] == "blocked"


def test_file_routes_reject_bad_ids(monkeypatch, tmp_path):
    client, app, _, root = make_client(monkeypatch, tmp_path)
    with client:
        r = client.post("/api/jobs/single", json={"prompt": OK_PROMPT, "steps": 4, "width": 512, "height": 512,
                                                   "image_name": "pic"})
        wait_job(client, r.json()["job_id"])
        good = client.get("/api/gallery/images").json()[0]["id"]
        assert client.get(f"/files/{good}.png").status_code == 200
        bad = ["/files/..%2f..%2fetc%2fpasswd", "/files/%2e%2e%2f%2e%2e%2fsecret.png",
               "/files/zzzzzzzzzzzzzzzz.png", "/files/ABCDEF0123456789.png", f"/files/{good}.PNG",
               f"/files/{good}x.png", f"/files/{good[:15]}.png", "/files/0123456789abcdef.png",
               "/files/not-hex-at-all/thumb.jpg", "/files/0123456789abcdef/thumb.jpg",
               "/api/images/..%2F..%2Fetc", "/api/images/0123456789abcdef", "/api/images/XYZ", "/nope"]
        for url in bad:
            r = client.get(url)
            assert r.status_code == 404, url
            assert set(r.json()["error"]) == {"cause", "fix"}
        for raw in (b"/files/..%2F", b"/files/..%2F..%2Fetc%2Fpasswd", b"/files/%2e%2e%2f", b"/files/..%5C..%5C",
                    b"/files/..%2F/thumb.jpg", b"/api/images/..%2F"):
            assert raw_status(app, raw) == 404, raw          # the exact bytes, no client-side normalising
        for bad_id in ("..%2F", "zzzz", "0123456789abcdef"):
            assert client.post(f"/api/images/{bad_id}/open-folder").status_code == 404
        assert client.post("/api/jobs/regenerate", json={"image_id": "../../x"}).status_code == 404
        assert client.post("/api/jobs/regenerate", json={"image_id": "0123456789abcdef"}).status_code == 404


def test_server_binds_loopback_only():
    src = SERVER_SRC.read_text(encoding="utf-8")
    hosts = re.findall(r"\bhost\s*=\s*([^,)\s]+)", src)
    assert hosts and all(h == '"127.0.0.1"' for h in hosts), hosts
    assert "0.0.0.0" not in src and "--host" not in src and '"::"' not in src


def test_foreign_host_and_cross_site_post_rejected(tmp_path):
    from fastapi.testclient import TestClient
    from medimg import config
    from medimg.engines.fake_engine import FakeEngine
    from medimg.server import create_app
    eng = FakeEngine(step_s=0.0)
    eng.load()
    s = config.Settings(output_root=tmp_path, unload_after_job=False)
    app = create_app(eng, s, None)
    base = f"http://127.0.0.1:{s.port}"
    with TestClient(app, base_url="http://evil.example:8750") as c:
        assert c.get("/api/config").status_code == 403
    with TestClient(app, base_url=base) as c:
        assert c.get("/api/config").status_code == 200
        assert c.post("/api/jobs/cancel").status_code == 403            # no X-Medimg
        r = c.post("/api/jobs/cancel", headers={"X-Medimg": "1", "Origin": "http://evil.example"})
        assert r.status_code == 403
