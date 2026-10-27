import hashlib
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent


def _load(name):
    spec = importlib.util.spec_from_file_location(name, REPO / "tools" / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


def _sha(p: Path):
    return hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None


def test_e2e_fake_engine_passes(tmp_path):
    e2e = _load("e2e_check")
    prof = REPO / "device_profile.json"
    before = _sha(prof)
    art = tmp_path / "art"
    r = subprocess.run([sys.executable, "tools/e2e_check.py", "--fake-engine", "--artifacts", str(art),
                        "--output-root", str(tmp_path / "out"), "--timeout", "180"],
                       cwd=REPO, capture_output=True, text=True, timeout=600)
    assert r.returncode == 0, r.stdout + r.stderr
    res = json.loads((art / "e2e_result.json").read_text(encoding="utf-8"))
    assert res["ok"] is True and res["schema"] == "medimg-e2e/1"
    assert [c["name"] for c in res["checks"]] == list(e2e.CHECK_NAMES)
    assert _sha(prof) == before


def test_prepare_sheet(tmp_path):
    e2e = _load("e2e_check")
    src = REPO / "templates" / "image_prompts_sample.xlsx"
    dst = tmp_path / "x.xlsx"
    e2e.prepare_sheet(src, dst, "512x512")
    a = load_workbook(src)["Image Prompts"]
    b = load_workbook(dst)["Image Prompts"]
    head = [c.value for c in a[1]]
    si, pi, ni = head.index("Size"), head.index("Image Prompt"), head.index("Image Name")
    assert b.max_row == a.max_row > 1
    for r in range(2, a.max_row + 1):
        assert b.cell(r, si + 1).value == "512x512"
        assert b.cell(r, pi + 1).value == a.cell(r, pi + 1).value
        assert b.cell(r, ni + 1).value == a.cell(r, ni + 1).value


def test_cache_key_format_and_sensitivity(tmp_path, monkeypatch):
    ck = _load("ci_cache_key")
    k = ck.cache_key()
    assert re.fullmatch(r"sdcpp-[0-9a-f]{16}", k)
    assert ck.cache_key() == k
    key_name = next(iter(ck.pins.ASSETS))
    asset = ck.pins.ASSETS[key_name]
    from dataclasses import replace
    monkeypatch.setitem(ck.pins.ASSETS, key_name, replace(asset, sha256="0" * 64))
    assert ck.cache_key() != k
    monkeypatch.undo()
    out = tmp_path / "gh_out"
    monkeypatch.setenv("GITHUB_OUTPUT", str(out))
    assert ck.main(["--tier", "low"]) == 0
    assert out.read_text(encoding="utf-8") == f"key={k}\n"


def test_profile_path_env_override(tmp_path):
    code = "from medimg import config; print(config.PROFILE_PATH)"
    import os
    env = {k: v for k, v in os.environ.items() if k != "MEDIMG_PROFILE_PATH"}
    a = subprocess.run([sys.executable, "-c", code], cwd=REPO, env=env, capture_output=True, text=True)
    assert Path(a.stdout.strip()) == REPO / "device_profile.json"
    target = tmp_path / "p.json"
    b = subprocess.run([sys.executable, "-c", code], cwd=REPO, env={**env, "MEDIMG_PROFILE_PATH": str(target)},
                       capture_output=True, text=True)
    assert Path(b.stdout.strip()) == target


def test_bad_sheet_path_exit_2(tmp_path):
    e2e = _load("e2e_check")
    assert e2e.main(["--fake-engine", "--sheet", str(tmp_path / "nope.xlsx"),
                     "--artifacts", str(tmp_path / "a")]) == 2
