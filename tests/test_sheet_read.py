import json

import pytest
from openpyxl import Workbook, load_workbook

from medimg.sheet import (REQUIRED, SheetError, current_sheet, find_worksheet, import_sheet,
                          read_rows)


def make(path, header, rows, title="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_header_matching_ignores_case_space_and_order(tmp_path):
    p = tmp_path / "a.xlsx"
    make(p, ["  image NAME ", "STATUS", "lecture", "Subject ", "unit", " Image   Prompt", "SEED", "size"],
         [["n1", None, "L", "S", "U", "a 45-year-old adult woman", 7, "768x1344"]])
    title, rows = read_rows(p)
    assert title == "Sheet1" and len(rows) == 1
    r = rows[0]
    assert (r.subject, r.unit, r.lecture, r.image_name) == ("S", "U", "L", "n1")
    assert r.prompt == "a 45-year-old adult woman"
    assert r.size == (768, 1344) and r.seed == 7 and r.idx == 2
    assert r.rel_path.endswith("n1.png") and r.rel_path.startswith("S/U/L/")


def test_blank_status_is_pending_and_blank_rows_skipped(tmp_path):
    p = tmp_path / "a.xlsx"
    make(p, ["Subject", "Unit", "Lecture", "Image Prompt", "Image Name", "Status"],
         [["S", "U", "L", "p", "n", None], [None] * 6, ["S", "U", "L", "p", "n2", "DONE"]])
    _, rows = read_rows(p)
    assert [r.status for r in rows] == ["pending", "done"]
    assert rows[1].idx == 4


def test_second_worksheet_found(tmp_path):
    p = tmp_path / "a.xlsx"
    wb = Workbook()
    wb.active.title = "Notes"
    wb.active.append(["nothing", "here"])
    ws = wb.create_sheet("Data")
    ws.append(["Subject", "Unit", "Lecture", "Image Prompt", "Image Name", "Status"])
    ws.append(["S", "U", "L", "p", "n", "pending"])
    wb.save(p)
    title, rows = read_rows(p)
    assert title == "Data" and len(rows) == 1
    assert find_worksheet(load_workbook(p)).title == "Data"


def test_no_matching_sheet_raises(tmp_path):
    p = tmp_path / "a.xlsx"
    make(p, ["Subject", "Unit"], [])
    with pytest.raises(SheetError) as e:
        read_rows(p)
    assert e.value.fix


def test_required_columns_pinned():
    assert REQUIRED == ("subject", "unit", "lecture", "image prompt", "image name", "status")


def test_bad_size_recorded(tmp_path):
    p = tmp_path / "a.xlsx"
    make(p, ["Subject", "Unit", "Lecture", "Image Prompt", "Image Name", "Status", "Size"],
         [["S", "U", "L", "p", "n", None, "big"]])
    _, rows = read_rows(p)
    assert rows[0].size is None and "768x1344" in rows[0].parse_error


def test_import_and_current(tmp_path):
    src = tmp_path / "in.xlsx"
    make(src, ["Subject"], [])
    root = tmp_path / "out"
    assert current_sheet(root) is None
    dest = import_sheet(src, root)
    assert dest.parent == root / "_sheets" and dest.name.endswith("_in.xlsx")
    assert dest.read_bytes() == src.read_bytes()
    assert json.loads((root / "_sheets" / "current.json").read_text())
    assert current_sheet(root) == dest
    dest.unlink()
    assert current_sheet(root) is None
