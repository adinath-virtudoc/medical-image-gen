import os

import pytest
from openpyxl import Workbook, load_workbook

from medimg.sheet import SYSTEM, ensure_system_columns, read_rows, write_row, write_rows

H = ["Notes", "Subject", "Unit", "Lecture", "Image Prompt", "Image Name", "Status"]


def build(path):
    wb = Workbook()
    other = wb.active
    other.title = "Other"
    other.append(["keep", "me"])
    ws = wb.create_sheet("Data")
    ws.append(H)
    for i in range(3):
        ws.append([f"note{i}", "S", "U", "L", "a 45-year-old adult", f"n{i}", None])
    ws.column_dimensions["B"].width = 33
    ws.column_dimensions["E"].width = 77
    wb.save(path)


def dump(path):
    wb = load_workbook(path)
    return {ws.title: [[c.value for c in r] for r in ws.iter_rows()] for ws in wb.worksheets}


def test_write_row_preserves_everything(tmp_path):
    p = tmp_path / "s.xlsx"
    build(p)
    before = dump(p)
    write_row(p, 3, {"Status": "done", "Seed Used": 42, "Output Path": "S/U/L/n1.png"})
    after = dump(p)
    assert after["Other"] == before["Other"]
    d0, d1 = before["Data"], after["Data"]
    assert [r[:7] for r in d1] == [r[:7] if i != 2 else r[:6] + ["done"] for i, r in enumerate(d0)]
    assert [r[0] for r in d1] == [r[0] for r in d0]            # row order intact
    wb = load_workbook(p)
    assert wb.sheetnames == ["Other", "Data"]
    assert wb["Data"].column_dimensions["B"].width == 33
    assert wb["Data"].column_dimensions["E"].width == 77
    hdr = [c.value for c in wb["Data"][1]]
    assert hdr[:7] == H and hdr[7:] == list(SYSTEM)
    assert wb["Data"].cell(row=3, column=hdr.index("Seed Used") + 1).value == 42


def test_ensure_system_columns_idempotent(tmp_path):
    p = tmp_path / "s.xlsx"
    build(p)
    m = ensure_system_columns(p)
    assert list(m)[-len(SYSTEM):] == list(SYSTEM)
    assert m["Steps"] == len(H) + 1 and m["Subject"] == 2
    assert ensure_system_columns(p) == m


def test_write_rows_many_and_reread(tmp_path):
    p = tmp_path / "s.xlsx"
    build(p)
    write_rows(p, {2: {"Status": "done", "Error": ""}, 4: {"Status": "failed", "Error": "boom"}})
    _, rows = read_rows(p)
    assert [r.status for r in rows] == ["done", "pending", "failed"]
    assert rows[2].error == "boom"


def test_atomic_write_survives_crash(tmp_path, monkeypatch):
    p = tmp_path / "s.xlsx"
    build(p)
    original = p.read_bytes()

    def boom(*a, **k):
        raise OSError("simulated crash")
    monkeypatch.setattr(os, "replace", boom)
    with pytest.raises(OSError):
        write_row(p, 2, {"Status": "done"})
    monkeypatch.undo()
    assert p.read_bytes() == original
    assert load_workbook(p)["Data"]["G2"].value is None


def test_formula_cells_frozen_to_cached_values(tmp_path):
    import re
    import zipfile

    p = tmp_path / "f.xlsx"
    build(p)
    wb = load_workbook(p)
    ws = wb["Data"]
    for r in (2, 3, 4):
        ws.cell(row=r, column=6).value = '="hand_"&ROW()'
    wb.save(p)
    # inject cached results like Excel would
    tmp = tmp_path / "g.xlsx"
    with zipfile.ZipFile(p) as zi, zipfile.ZipFile(tmp, "w") as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "xl/worksheets/sheet2.xml":
                txt = data.decode()
                n = iter((2, 3, 4))
                txt = re.sub(r'<c r="(F\d+)"[^>]*>(<f>.*?</f>)(?:<v\s*/>|<v></v>)',
                             lambda m: f'<c r="{m.group(1)}" t="str">' + m.group(2) + f"<v>hand_{next(n)}</v>", txt)
                assert "<v>hand_4</v>" in txt
                data = txt.encode()
            zo.writestr(item, data)
    os.replace(tmp, p)
    assert [r.image_name for r in read_rows(p)[1]] == ["hand_2", "hand_3", "hand_4"]
    write_row(p, 2, {"Status": "done"})
    assert [r.image_name for r in read_rows(p)[1]] == ["hand_2", "hand_3", "hand_4"]
