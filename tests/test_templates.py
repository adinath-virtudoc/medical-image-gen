import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook

from medimg.sheet import REQUIRED, find_worksheet, read_rows

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "templates"


def test_generator_runs_into_tmp(tmp_path):
    spec = importlib.util.spec_from_file_location("make_templates", ROOT / "tools" / "make_templates.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.main(tmp_path) == 0
    assert (tmp_path / "image_prompts_template.xlsx").is_file()
    assert (tmp_path / "image_prompts_sample.xlsx").is_file()


def test_template_has_status_validation_size_comment_and_help():
    wb = load_workbook(TEMPLATES / "image_prompts_template.xlsx")
    ws = find_worksheet(wb)
    assert [str(c.value).lower() for c in ws[1]][:6] == list(REQUIRED[:3]) + ["image prompt", "image name", "status"]
    assert ws.max_row == 1
    dvs = ws.data_validations.dataValidation
    assert any(dv.type == "list" and "F2" in str(dv.sqref) and
               all(s in dv.formula1 for s in ("pending", "done", "failed", "blocked", "invalid"))
               for dv in dvs)
    assert "768x1344" in ws["G1"].comment.text
    assert "How to fill" in wb.sheetnames


def test_sample_has_three_adult_rows_that_pass_rules():
    title, rows = read_rows(TEMPLATES / "image_prompts_sample.xlsx")
    assert len(rows) == 3 and all(r.status == "pending" for r in rows)
    check = pytest.importorskip("medimg.rules").check_prompt
    assert [check(r.prompt) for r in rows] == [None, None, None]
