import threading
import time

import pytest
from openpyxl import Workbook, load_workbook

from medimg.assess import fit, load_profile, save_profile
from medimg.engines.fake_engine import FakeEngine
from medimg.errors import Blocked, Busy, ModelUnavailable
from medimg.jobs import JobSpec, RowJob
from medimg.store import Store
from medimg.worker import ACTIVE, GenerationWorker

HEAD = ["Subject", "Unit", "Lecture", "Image Prompt", "Image Name", "Status"]
PROMPT = "a 45-year-old adult man, clinic photo"


def make_sheet(path, n=3):
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD)
    for i in range(n):
        ws.append(["Anat", "Hand", "L1", PROMPT, f"img{i}", "pending"])
    wb.save(path)


def sheet_rows(path):
    ws = load_workbook(path).worksheets[0]
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


def profile_dict():
    pts = [{"width": 512, "height": 512, "steps": 4, "total_s": 30.0, "per_step_s": 5.5, "overhead_s": 8.0},
           {"width": 1024, "height": 1024, "steps": 2, "total_s": 62.0, "per_step_s": 23.0, "overhead_s": 16.0}]
    return {"schema_version": 1, "created_at": "x", "app_version": "0", "backend": {},
            "machine": {"hostname": "testhost", "apple_chip": "Apple Test"},
            "benchmark": {"load_s": 1.0, "points": pts, "fit": fit(pts)}, "actuals": [], "notes": []}


def rowjobs(n=3, **kw):
    return [RowJob(prompt=PROMPT, seed=None, steps=4, width=512, height=512, subject="Anat",
                   unit="Hand", lecture="L1", image_name=f"img{i}", sheet_row=i + 2, **kw)
            for i in range(n)]


def wait_idle(w, timeout=15):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        if w.snapshot()["job"] is None and w.snapshot()["last_job"]:
            return w.snapshot()
        time.sleep(0.01)
    raise AssertionError("worker did not finish: %r" % (w.snapshot(),))


@pytest.fixture
def env(tmp_path):
    store = Store(tmp_path / "out")
    sheet_path = tmp_path / "work.xlsx"
    make_sheet(sheet_path)
    pp = tmp_path / "device_profile.json"
    save_profile(profile_dict(), pp)
    made = []

    def build(engine=None, unload=False):
        w = GenerationWorker(engine or FakeEngine(step_s=0.0), store, pp, unload_after_job=unload,
                             sheet_lock=threading.Lock())
        w.start()
        made.append(w)
        return w

    yield build, store, sheet_path, pp
    for w in made:
        w.stop()


def batch(sheet_path, rows):
    return JobSpec(kind="batch", rows=rows, sheet_path=sheet_path, sheet_name="orig.xlsx")


def test_batch_three_rows(env):
    build, store, sp, pp = env
    w = build()
    jid = w.submit(batch(sp, rowjobs()))
    snap = wait_idle(w)
    assert snap["last_job"]["id"] == jid and snap["last_job"]["outcome"] == "done"
    assert snap["model"]["state"] == "ready"
    rows = sheet_rows(sp)
    for i, r in enumerate(rows):
        assert r["Status"] == "done"
        assert r["Seed Used"] is not None and r["Seed"] == r["Seed Used"]
        assert r["Steps"] == 4 and r["Model"] == "fake/fake-model"
        assert r["Machine"] == "testhost · Apple Test"
        assert r["Output Path"] == f"Anat/Hand/L1/img{i}.png"
        assert r["Generated At"][-6] in "+-" and r["Error"] is None
        assert (store.root / r["Output Path"]).exists()
        assert (store.root / r["Output Path"]).with_suffix(".json").exists()
    assert len(store.list_images("Anat", "Hand", "L1")) == 3
    assert len(load_profile(pp)["actuals"]) == 3
    assert snap["gallery_version"] == 3 and snap["version"] > 0


def test_snapshot_shape_while_running(env):
    build, store, sp, pp = env
    w = build(FakeEngine(step_s=0.05))
    w.submit(batch(sp, rowjobs(2)))
    seen = None
    end = time.monotonic() + 10
    while time.monotonic() < end:
        s = w.snapshot()
        if s["job"] and s["job"]["phase"] == "denoising":
            seen = s
            break
        time.sleep(0.005)
    assert seen and seen["job"]["phase"] in ACTIVE
    assert set(seen["job"]) == {"id", "kind", "phase", "step", "step_total", "image_index",
                                "image_count", "eta_s", "current", "done_ids", "errors"}
    assert seen["job"]["image_count"] == 2 and seen["job"]["eta_s"] > 0
    assert seen["job"]["current"]["image_name"] == "img0"
    assert set(seen["model"]) == {"state", "error"} and seen["local_only"] is True
    w.cancel()
    wait_idle(w)


def test_cancel_after_row_one(env):
    build, store, sp, pp = env

    class Cancelling(FakeEngine):
        def generate(self, *a, **kw):
            if self.calls == 1:           # entering the 2nd call
                worker.cancel()
            return super().generate(*a, **kw)

    worker = build(Cancelling(step_s=0.0))
    worker.submit(batch(sp, rowjobs()))
    snap = wait_idle(worker)
    assert snap["last_job"]["outcome"] == "cancelled"
    r = sheet_rows(sp)
    assert [x["Status"] for x in r] == ["done", "pending", "pending"]
    assert (store.root / "Anat/Hand/L1/img0.png").exists()
    assert not (store.root / "Anat/Hand/L1/img1.png").exists()
    assert not (store.root / "Anat/Hand/L1/img2.png").exists()
    assert r[1].get("Seed Used") is None and r[1].get("Output Path") is None
    assert worker.cancel() is False


def test_row_failure_continues(env):
    build, store, sp, pp = env
    w = build(FakeEngine(step_s=0.0, fail_on={2}))
    w.submit(batch(sp, rowjobs()))
    snap = wait_idle(w)
    r = sheet_rows(sp)
    assert [x["Status"] for x in r] == ["done", "failed", "done"]
    assert "fake failure" in r[1]["Error"]
    assert snap["last_job"]["outcome"] == "done"
    assert not (store.root / "Anat/Hand/L1/img1.png").exists()


def test_blocked_row_in_batch_written_not_generated(env):
    build, store, sp, pp = env
    rows = rowjobs(2)
    rows[0].prompt = "a teen boy at the clinic"
    eng = FakeEngine(step_s=0.0)
    w = build(eng)
    w.submit(batch(sp, rows))
    wait_idle(w)
    r = sheet_rows(sp)
    assert r[0]["Status"] == "blocked" and r[0]["Error"].startswith("Blocked: ")
    assert r[1]["Status"] == "done"
    assert eng.calls == 1


def test_blocked_prompt_raises_on_direct_submit(env):
    build, store, sp, pp = env
    w = build()
    row = RowJob(prompt="a teenager", seed=1, steps=4, width=512, height=512, subject="", unit="",
                 lecture="", image_name="x")
    with pytest.raises(Blocked):
        w.submit(JobSpec(kind="single", rows=[row]))
    assert w.snapshot()["job"] is None


def test_second_submit_busy(env):
    build, store, sp, pp = env
    w = build(FakeEngine(step_s=0.05))
    w.submit(batch(sp, rowjobs(2)))
    with pytest.raises(Busy):
        w.submit(batch(sp, rowjobs(1)))
    w.cancel()
    wait_idle(w)


def test_single_goes_to_single_folder_no_sheet(env):
    build, store, sp, pp = env
    w = build()
    row = RowJob(prompt=PROMPT, seed=7, steps=4, width=770, height=530, subject="", unit="",
                 lecture="", image_name="solo")
    w.submit(JobSpec(kind="single", rows=[row]))
    wait_idle(w)
    rec = store.list_images("_single", None, None)[0]
    assert rec.rel_path == "_single/solo.png" and rec.seed == 7 and rec.source == "single"
    assert (rec.width, rec.height) == (768, 528)


def test_regenerate_overwrites_and_updates_row(env):
    build, store, sp, pp = env
    w = build()
    w.submit(batch(sp, rowjobs(1)))
    wait_idle(w)
    first = sheet_rows(sp)[0]["Seed Used"]
    rows = rowjobs(1, overwrite=True)
    rows[0].seed = first + 1
    w.submit(JobSpec(kind="regenerate", rows=rows, sheet_path=sp, sheet_name="orig.xlsx"))
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and sheet_rows(sp)[0]["Seed Used"] == first:
        time.sleep(0.01)
    wait_idle(w)
    r = sheet_rows(sp)[0]
    assert r["Seed Used"] == first + 1 and r["Seed"] == first + 1 and r["Status"] == "done"
    assert store.list_images("Anat", "Hand", "L1")[0].source == "regenerate"


def test_existing_file_without_overwrite_fails_row(env):
    build, store, sp, pp = env
    w = build()
    w.submit(batch(sp, rowjobs(1)))
    wait_idle(w)
    w.submit(batch(sp, rowjobs(1)))
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and w.snapshot()["job"] is not None:
        time.sleep(0.01)
    assert sheet_rows(sp)[0]["Status"] == "failed"


def test_unload_after_job(env):
    build, store, sp, pp = env

    class Counting(FakeEngine):
        unloads = 0

        def unload(self):
            Counting.unloads += 1
            super().unload()

    eng = Counting(step_s=0.0)
    w = build(eng, unload=True)
    w.submit(batch(sp, rowjobs(2)))
    snap = wait_idle(w)
    assert Counting.unloads == 1 and eng.loaded is False
    assert snap["model"]["state"] == "not_loaded"
    sp2 = sp  # second job reloads the model
    w.submit(batch(sp2, rowjobs(1, overwrite=True)))
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and Counting.unloads < 2:
        time.sleep(0.01)
    assert Counting.unloads == 2


def test_load_error(env):
    build, store, sp, pp = env

    class Broken(FakeEngine):
        def load(self, on_progress=None):
            raise RuntimeError("weights missing")

    w = build(Broken(step_s=0.0))
    w.submit(batch(sp, rowjobs()))
    snap = wait_idle(w)
    assert snap["last_job"]["outcome"] == "failed"
    assert snap["model"]["state"] == "error" and "weights missing" in snap["model"]["error"]
    assert [x["Status"] for x in sheet_rows(sp)] == ["pending"] * 3
    with pytest.raises(ModelUnavailable):
        w.submit(batch(sp, rowjobs(1)))
    w.clear_model_error()
    assert w.snapshot()["model"]["state"] == "not_loaded"


def test_benchmark_kind_saves_profile(env):
    build, store, sp, pp = env
    w = build()
    w.submit(JobSpec(kind="benchmark"))
    snap = wait_idle(w)
    assert snap["last_job"]["kind"] == "benchmark" and snap["last_job"]["outcome"] == "done"
    prof = load_profile(pp)
    assert len(prof["benchmark"]["points"]) == 2 and "fit" in prof["benchmark"]


def test_worker_does_not_write_foreign_profile(env):
    build, store, sp, pp = env

    class SdFake(FakeEngine):
        name = "sdcpp"

    prof = profile_dict()
    prof["backend"] = {"name": "mflux"}
    save_profile(prof, pp)
    before = pp.read_bytes()
    w = build(SdFake(step_s=0.0))
    w.submit(batch(sp, rowjobs(2)))
    wait_idle(w)
    assert pp.read_bytes() == before
    assert load_profile(pp)["actuals"] == []

    prof["backend"] = {"name": "sdcpp"}
    save_profile(prof, pp)
    w2 = build(SdFake(step_s=0.0))
    sp2 = sp.parent / "work2.xlsx"
    make_sheet(sp2)
    rows2 = rowjobs(1)
    rows2[0].image_name = "other_img"
    w2.submit(batch(sp2, rows2))
    wait_idle(w2)
    assert len(load_profile(pp)["actuals"]) == 1
