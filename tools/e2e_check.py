"""End-to-end check: start the real server, drive it over HTTP, verify the result.

    python tools/e2e_check.py [--engine {auto,sdcpp}] [--fake-engine] [--port N]
                              [--output-root DIR] [--artifacts DIR] [--size WxH]
                              [--steps N] [--timeout SECONDS] [--sheet PATH]

Exit 0 = every check passed; 1 = a check failed; 2 = the server never started or bad arguments.
Uses only the standard library, openpyxl and Pillow. The server only ever binds 127.0.0.1.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import uuid
from datetime import datetime, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
HOST = "127.0.0.1"

CHECK_NAMES = (
    "server_started", "config_local_only", "sheet_uploaded", "gate_blocks_minor", "batch_submitted",
    "job_finished", "gallery_lists_image", "png_written", "sidecar_written", "sheet_row_done",
)
PNG_SIG = b"\x89PNG\r\n\x1a\n"


def prepare_sheet(src: Path, dst: Path, size: str) -> None:
    """Copy the workbook to dst and set every data row's Size cell to size."""
    from openpyxl import load_workbook
    wb = load_workbook(src)
    ws = wb["Image Prompts"] if "Image Prompts" in wb.sheetnames else wb.worksheets[0]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = header.index("Size") + 1
    for r in range(2, ws.max_row + 1):
        if any(c.value not in (None, "") for c in ws[r]):
            ws.cell(row=r, column=col).value = size
    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)


def _free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind((HOST, 0))
        return s.getsockname()[1]


def _http(method: str, url: str, *, data: bytes | None = None, headers: dict | None = None,
          timeout: float = 30.0) -> tuple[int, bytes]:
    h = dict(headers or {})
    if method == "POST":
        h["X-Medimg"] = "1"
    req = urllib.request.Request(url, data=data, method=method, headers=h)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.status, resp.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def _json(method: str, url: str, body: dict | None = None, timeout: float = 30.0) -> tuple[int, object]:
    data = json.dumps(body).encode("utf-8") if body is not None else None
    hdr = {"Content-Type": "application/json"} if body is not None else {}
    code, raw = _http(method, url, data=data, headers=hdr, timeout=timeout)
    try:
        return code, json.loads(raw.decode("utf-8"))
    except ValueError:
        return code, {"raw": raw[:200].decode("utf-8", "replace")}


def _multipart(filename: str, content: bytes) -> tuple[bytes, str]:
    boundary = "----medimg" + uuid.uuid4().hex
    parts = [
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
        "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n",
    ]
    body = parts[0].encode("utf-8") + content + f"\r\n--{boundary}--\r\n".encode("utf-8")
    return body, f"multipart/form-data; boundary={boundary}"


def _stop(proc: subprocess.Popen) -> None:
    if proc.poll() is not None:
        return
    try:
        if os.name == "nt":
            proc.terminate()
            proc.wait(timeout=15)
        else:
            proc.send_signal(signal.SIGINT)
            proc.wait(timeout=15)
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.wait()
    except OSError:
        pass


class _Run:
    def __init__(self) -> None:
        self.checks: list[dict] = []

    def record(self, name: str, ok: bool, detail: str) -> bool:
        self.checks.append({"name": name, "ok": bool(ok), "detail": detail})
        print(f"{'PASS' if ok else 'FAIL'}  {name}: {detail}", flush=True)
        return bool(ok)

    def skip_rest(self, why: str) -> None:
        done = {c["name"] for c in self.checks}
        for n in CHECK_NAMES:
            if n not in done:
                self.record(n, False, f"skipped: {why}")


def _parse_args(argv):
    ap = argparse.ArgumentParser(description="End-to-end check of the local server.")
    ap.add_argument("--engine", choices=("auto", "sdcpp"), default=None)
    ap.add_argument("--fake-engine", action="store_true")
    ap.add_argument("--port", type=int, default=0, help="0 = pick a free port")
    ap.add_argument("--output-root", default=None)
    ap.add_argument("--artifacts", default="e2e-artifacts")
    ap.add_argument("--size", default="512x512")
    ap.add_argument("--steps", type=int, default=4)
    ap.add_argument("--timeout", type=int, default=3600)
    ap.add_argument("--sheet", default=str(REPO_ROOT / "templates" / "image_prompts_sample.xlsx"))
    return ap.parse_args(argv)


def main(argv=None) -> int:
    try:
        args = _parse_args(argv)
    except SystemExit as e:
        return 2 if e.code else 0
    sheet = Path(args.sheet).expanduser()
    if not sheet.is_file():
        print(f"error: sheet not found: {sheet}", file=sys.stderr)
        return 2
    import re
    if not re.fullmatch(r"\d+x\d+", args.size):
        print(f"error: --size must look like 512x512, got {args.size!r}", file=sys.stderr)
        return 2
    if not (0 <= args.port <= 65535):
        print("error: --port must be 0-65535", file=sys.stderr)
        return 2
    width, height = (int(x) for x in args.size.split("x"))

    art = Path(args.artifacts).expanduser().resolve()
    art.mkdir(parents=True, exist_ok=True)
    out_root = Path(args.output_root).expanduser().resolve() if args.output_root \
        else Path(tempfile.mkdtemp(prefix="medimg-e2e-"))
    out_root.mkdir(parents=True, exist_ok=True)
    port = args.port or _free_port()
    base = f"http://{HOST}:{port}"

    started = datetime.now(timezone.utc)
    t0 = time.monotonic()
    run = _Run()
    png_name = None
    device = None
    cfg: dict = {}
    sheet_up = art / "sheet_uploaded.xlsx"
    prepare_sheet(sheet, sheet_up, args.size)

    cmd = [sys.executable, "-m", "medimg", "serve", "--no-browser", "--port", str(port),
           "--output-root", str(out_root)]
    env = dict(os.environ, PYTHONUTF8="1")
    if args.fake_engine:
        cmd.append("--fake-engine")
        env["MEDIMG_PROFILE_PATH"] = str(art / "device_profile.json")
    elif args.engine:
        cmd += ["--engine", args.engine]
    log = open(art / "server.log", "wb")
    proc = subprocess.Popen(cmd, cwd=str(REPO_ROOT), env=env, stdout=log, stderr=subprocess.STDOUT)
    started_ok = False
    try:
        # 1. server_started
        deadline = time.monotonic() + 120
        last = "no answer"
        while time.monotonic() < deadline:
            if proc.poll() is not None:
                last = f"server exited with code {proc.returncode} (see server.log)"
                break
            try:
                code, body = _json("GET", base + "/api/config", timeout=5)
                if code == 200 and isinstance(body, dict):
                    cfg, started_ok = body, True
                    break
                last = f"HTTP {code}"
            except (urllib.error.URLError, OSError) as e:
                last = str(e)
            time.sleep(0.5)
        run.record("server_started", started_ok,
                   f"{base} up, backend {cfg.get('backend')}" if started_ok else last)
        if not started_ok:
            run.skip_rest("server did not start")
            return _finish(run, args, art, t0, started, None, None, 2)
        device = cfg.get("device")

        # 2. config_local_only
        run.record("config_local_only", cfg.get("local_only") is True, f"local_only={cfg.get('local_only')}")

        # 3. sheet_uploaded
        body, ctype = _multipart("e2e_sheet.xlsx", sheet_up.read_bytes())
        code, rep = _http("POST", base + "/api/sheet", data=body, headers={"Content-Type": ctype}, timeout=60)
        try:
            rep = json.loads(rep.decode("utf-8"))
        except ValueError:
            rep = {}
        counts = rep.get("counts", {}) if isinstance(rep, dict) else {}
        run.record("sheet_uploaded", code == 200 and counts.get("total") == 3 and counts.get("pending") == 3,
                   f"HTTP {code}, counts {counts}")

        # 4. gate_blocks_minor
        code, err = _json("POST", base + "/api/jobs/single",
                          {"prompt": "Portrait photo of a 12-year-old child"})
        cause = (err.get("error", {}) if isinstance(err, dict) else {}).get("cause", "")
        run.record("gate_blocks_minor", code == 422 and str(cause).startswith("Blocked"),
                   f"HTTP {code}, cause {str(cause)[:80]!r}")

        # 5. batch_submitted
        code, sub = _json("POST", base + "/api/jobs/batch",
                          {"count": 1, "default_size": args.size, "steps": args.steps})
        batch_ok = run.record("batch_submitted", code == 202, f"HTTP {code} {sub if code != 202 else ''}".strip())

        # 6. job_finished
        job_id = sub.get("job_id") if isinstance(sub, dict) else None
        finished = False
        if batch_ok:
            deadline = time.monotonic() + args.timeout
            errors: list = []
            outcome = None
            seen_job = False
            detail = "timed out"
            while time.monotonic() < deadline:
                try:
                    _c, st = _json("GET", base + "/api/status", timeout=15)
                except (urllib.error.URLError, OSError) as e:
                    detail = f"status request failed: {e}"
                    time.sleep(2)
                    continue
                job = st.get("job") if isinstance(st, dict) else None
                if job:
                    seen_job = True
                    for e in job.get("errors") or []:
                        if e not in errors:
                            errors.append(e)
                last_job = (st.get("last_job") if isinstance(st, dict) else None) or {}
                if job is None and last_job and (job_id is None or last_job.get("id") == job_id):
                    outcome = last_job.get("outcome")
                    break
                if proc.poll() is not None:
                    detail = f"server exited with code {proc.returncode}"
                    break
                time.sleep(2)
            if outcome is not None:
                finished = outcome == "done" and not errors
                detail = f"outcome {outcome}" + (f", job errors {errors}" if errors else "")
            run.record("job_finished", finished, detail)
        else:
            run.record("job_finished", False, "skipped: batch was not submitted")

        # 7. gallery_lists_image
        imgs: list = []
        try:
            _c, imgs = _json("GET", base + "/api/gallery/images")
        except (urllib.error.URLError, OSError):
            imgs = []
        imgs = imgs if isinstance(imgs, list) else []
        run.record("gallery_lists_image", len(imgs) == 1, f"{len(imgs)} image(s) listed")
        rec = imgs[0] if imgs else None
        rel = rec.get("rel_path") if rec else None
        image_name = rec.get("image_name") if rec else None

        # 8. png_written
        png_path = None
        if rel:
            png_path = out_root / rel
            ok = False
            detail = f"{png_path} missing"
            if png_path.is_file():
                data = png_path.read_bytes()
                if not data.startswith(PNG_SIG):
                    detail = "file is not a PNG"
                else:
                    from PIL import Image
                    with Image.open(png_path) as im:
                        got = im.size
                    ok = args.fake_engine or got == (width, height)
                    detail = f"{rel} PNG {got[0]}x{got[1]}" + ("" if ok else f", expected {width}x{height}")
            run.record("png_written", ok, detail)
            if png_path.is_file():
                shutil.copy2(png_path, art / png_path.name)
                png_name = png_path.name
        else:
            run.record("png_written", False, "no image to check")

        # 9. sidecar_written
        side = png_path.with_suffix(".json") if png_path else None
        if side and side.is_file():
            try:
                meta = json.loads(side.read_text(encoding="utf-8"))
            except ValueError:
                meta = {}
            missing = [k for k in ("engine", "model_id", "seed", "steps", "time_s") if k not in meta]
            run.record("sidecar_written", not missing, f"{side.name}" + (f", missing {missing}" if missing else ""))
            shutil.copy2(side, art / side.name)
        else:
            run.record("sidecar_written", False, "sidecar missing")

        # 10. sheet_row_done
        ok, detail = False, "no image to match"
        try:
            code, raw = _http("GET", base + "/api/sheet/download", timeout=60)
            if code == 200:
                after = art / "sheet_after.xlsx"
                after.write_bytes(raw)
                from openpyxl import load_workbook
                wb = load_workbook(after, data_only=True)
                ws = wb["Image Prompts"] if "Image Prompts" in wb.sheetnames else wb.worksheets[0]
                head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
                ci, cs, co = head.index("Image Name"), head.index("Status"), head.index("Output Path")
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if image_name and row[ci] == image_name:
                        ok = row[cs] == "done" and row[co] == rel
                        detail = f"row {image_name}: Status {row[cs]!r}, Output Path {row[co]!r}"
                        break
                else:
                    detail = f"no row named {image_name!r}"
            else:
                detail = f"download HTTP {code}"
        except Exception as e:  # keep the summary even if the workbook is odd
            detail = f"{type(e).__name__}: {e}"
        run.record("sheet_row_done", ok, detail)
    finally:
        _stop(proc)
        log.close()
    all_ok = all(c["ok"] for c in run.checks) and len(run.checks) == len(CHECK_NAMES)
    return _finish(run, args, art, t0, started, device, png_name, 0 if all_ok else 1)


def _finish(run: _Run, args, art: Path, t0: float, started: datetime, device, png, code: int) -> int:
    result = {
        "schema": "medimg-e2e/1", "ok": code == 0, "checks": run.checks,
        "engine": "fake" if args.fake_engine else (args.engine or "auto"), "device": device,
        "seconds": round(time.monotonic() - t0, 1), "png": png,
        "started_at": started.isoformat(), "finished_at": datetime.now(timezone.utc).isoformat(),
    }
    (art / "e2e_result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    passed = sum(1 for c in run.checks if c["ok"])
    print(f"\n{'OK' if code == 0 else 'FAILED'}: {passed}/{len(CHECK_NAMES)} checks passed "
          f"in {result['seconds']} s. Details: {art / 'e2e_result.json'}", flush=True)
    return code


if __name__ == "__main__":
    sys.exit(main())
