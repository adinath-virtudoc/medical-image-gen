"""Generate templates/image_prompts_template.xlsx and templates/image_prompts_sample.xlsx."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["Subject", "Unit", "Lecture", "Image Prompt", "Image Name", "Status", "Size", "Seed"]
WIDTHS = [16, 18, 18, 70, 24, 12, 14, 12]
STATUS_LIST = "pending,done,failed,blocked,invalid"

SAMPLE_ROWS = [
    ["Anatomy", "Upper limb", "Hand",
     "X-ray-style illustration of the bones of the right hand of a 40-year-old adult, "
     "clean white background, labelled-diagram look, no text", "hand_xray_adult", "pending", "1024x1024", None],
    ["Anatomy", "Lower limb", "Knee",
     "Photo of a 52-year-old adult woman in a clinic, wearing shorts and a plain t-shirt, seated "
     "while a clinician examines her knee, neutral lighting", "knee_exam_adult", "pending", "1344x768", None],
    ["Anatomy", "Upper limb", "Shoulder",
     "Clear anatomical diagram of the shoulder joint of a 35-year-old adult man wearing a fitted "
     "sleeveless top, muscles and tendons outlined, plain background", "shoulder_diagram_adult", "pending", "768x1344", None],
]

HOW_TO_FILL = [
    ("How to fill this sheet", ""),
    ("", ""),
    ("Subject / Unit / Lecture", "Required. Each becomes a folder: Subject/Unit/Lecture/."),
    ("Image Prompt", "Required. Describe the image. Always state an adult age (for example "
                     "45-year-old) or the word adult. Prompts that suggest a minor are blocked."),
    ("Image Name", "Required. File name without extension. The tool adds .png."),
    ("Status", "Leave blank (= pending). The tool writes done, failed, blocked or invalid. "
               "Done rows are never regenerated."),
    ("Size", "Optional. WxH, for example 768x1344. Blank uses the size chosen in the app."),
    ("Seed", "Optional. Blank gets a random seed, which is written back."),
    ("System columns", "Steps, Seed Used, Generated At, Machine, Model, Generation Time (s), "
                       "Output Path and Error are added and filled by the tool. Do not edit them."),
    ("Column order", "Does not matter. Header case and spaces are ignored."),
]


def _header_row(ws) -> None:
    ws.append(HEADERS)
    for i, (cell, w) in enumerate(zip(ws[1], WIDTHS), start=1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws["G1"].comment = Comment("WxH, e.g. 768x1344; blank = UI default", "medical-image-gen")
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1=f'"{STATUS_LIST}"', allow_blank=True)
    dv.error, dv.errorTitle = "Pick a status from the list.", "Invalid status"
    ws.add_data_validation(dv)
    dv.add("F2:F1000")


def _how_to(wb) -> None:
    ws = wb.create_sheet("How to fill")
    for a, b in HOW_TO_FILL:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=14)
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).font = Font(bold=True)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90


def make_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Image Prompts"
    _header_row(ws)
    _how_to(wb)
    wb.save(path)


def make_sample(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Image Prompts"
    _header_row(ws)
    for row in SAMPLE_ROWS:
        ws.append(row)
        ws.cell(row=ws.max_row, column=4).alignment = Alignment(wrap_text=True, vertical="top")
    _how_to(wb)
    wb.save(path)


def main(out_dir: Path | None = None) -> int:
    out = Path(out_dir) if out_dir else Path(__file__).resolve().parent.parent / "templates"
    out.mkdir(parents=True, exist_ok=True)
    make_template(out / "image_prompts_template.xlsx")
    make_sample(out / "image_prompts_sample.xlsx")
    print(f"Wrote templates to {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
